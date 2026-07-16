"""
DOCX文档解析器

使用 python-docx 提取文本、段落、标题、表格等结构信息。

成员5主责：DOCX解析实现
"""
import io
from typing import BinaryIO

from app.core.exceptions import DocumentProcessingError
from app.services.parsing.base import (
    BaseParser,
    ParseResult,
    PageElement,
    TableElement,
    ParseQuality,
)


class DOCXParser(BaseParser):
    """DOCX文档解析器"""

    @property
    def supported_format(self) -> str:
        return "docx"

    def parse(self, data: bytes, filename: str = "") -> ParseResult:
        """解析DOCX文件"""
        result = ParseResult(file_type="docx")
        file_obj = io.BytesIO(data)

        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(file_obj)
            page_element = PageElement(page_number=1)
            position = 0
            full_text_parts = []

            # 遍历文档体元素
            for element in doc.element.body:
                tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

                if tag == "p":
                    # 段落处理
                    para_text = self._extract_paragraph_text(element, doc)
                    if para_text:
                        # 判断段落样式
                        style = self._get_paragraph_style(element, doc)
                        if style and style.startswith("Heading"):
                            level = int(style.replace("Heading", "")) if style != "Heading" else 1
                            page_element.headings.append({
                                "level": min(level, 6),
                                "text": para_text,
                                "position": position,
                            })
                        else:
                            page_element.paragraphs.append(para_text)

                        full_text_parts.append(para_text)
                        position += len(para_text) + 1

                elif tag == "tbl":
                    # 表格处理
                    table_elem = self._extract_table(element)
                    if table_elem:
                        page_element.tables.append(table_elem)
                        result.tables.append(table_elem)
                        # 表格文本加入全文
                        for row in table_elem.rows:
                            full_text_parts.append(" | ".join(row))
                            position += len(" | ".join(row)) + 1

            page_element.text = "\n".join(full_text_parts)
            result.total_chars = len(page_element.text)
            result.total_pages = 1
            result.pages.append(page_element)

        except Exception as e:
            raise DocumentProcessingError(
                message=f"DOCX解析失败: {str(e)}",
                details={"filename": filename, "error": str(e)},
            )

        return result

    def parse_stream(self, stream: BinaryIO, filename: str = "") -> ParseResult:
        """流式解析DOCX"""
        data = stream.read()
        return self.parse(data, filename)

    def _extract_paragraph_text(self, element, doc) -> str:
        """提取段落文本"""
        texts = []
        for child in element.iter():
            if child.tag.endswith("}t"):
                if child.text:
                    texts.append(child.text)
        return "".join(texts).strip()

    def _get_paragraph_style(self, element, doc) -> str | None:
        """获取段落样式名称"""
        for child in element.iter():
            style_elem = child.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
            if style_elem is not None:
                style_id = style_elem.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                if style_id:
                    try:
                        style = doc.styles[style_id]
                        return style.name if style else None
                    except KeyError:
                        return style_id
        return None

    def _extract_table(self, element) -> TableElement | None:
        """提取表格数据"""
        rows_data = []
        for row_elem in element.iter():
            if row_elem.tag.endswith("}tr"):
                cells = []
                for cell_elem in row_elem.iter():
                    if cell_elem.tag.endswith("}tc"):
                        cell_texts = []
                        for para in cell_elem.iter():
                            if para.tag.endswith("}t") and para.text:
                                cell_texts.append(para.text)
                        cells.append("".join(cell_texts).strip())
                if cells:
                    rows_data.append(cells)

        if not rows_data or len(rows_data) < 1:
            return None

        headers = rows_data[0] if len(rows_data) > 1 else []
        data_rows = rows_data[1:] if len(rows_data) > 1 else rows_data

        return TableElement(headers=headers, rows=data_rows)
